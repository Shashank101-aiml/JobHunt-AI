import pytest
import os
import tempfile
import docx

from app.services.resume_parser import (
    extract_contact_info,
    extract_skills_nlp,
    extract_sections,
    parse_resume_document,
    extract_text_by_type
)

SAMPLE_RESUME_TEXT = """Jane Doe
jane.doe@techcorp.io | +1 (415) 555-0199 | Austin, TX
linkedin.com/in/janedoe-eng | github.com/janedoe

PROFESSIONAL SUMMARY
Principal Software Architect with 8+ years building enterprise distributed platforms.

TECHNICAL SKILLS
Python, Go, FastAPI, Docker, Kubernetes, PostgreSQL, Redis, AWS, Terraform

WORK EXPERIENCE
Senior Lead Engineer | CloudData Systems | 2021 - Present
- Architected event-driven microservices using FastAPI and Kafka.
- Reduced database bottleneck using Redis caching and PostgreSQL query optimization.

Software Engineer | AppForge | 2017 - 2021
- Built REST APIs in Python and Django with full CI/CD automation.

EDUCATION
Bachelor of Science in Computer Science | University of Texas | 2017
"""

def test_extract_contact_info():
    info = extract_contact_info(SAMPLE_RESUME_TEXT)
    assert info.email == "jane.doe@techcorp.io"
    assert info.name == "Jane Doe"
    assert "415" in info.phone
    assert "linkedin.com/in/janedoe-eng" in info.linkedin
    assert "github.com/janedoe" in info.github

def test_extract_skills_nlp():
    skills = extract_skills_nlp(SAMPLE_RESUME_TEXT)
    assert "Python" in skills
    assert "FastAPI" in skills
    assert "Docker" in skills
    assert "Kubernetes" in skills
    assert "PostgreSQL" in skills
    assert "AWS" in skills
    # Skills not in text should not be present
    assert "Ruby" not in skills
    assert "PHP" not in skills

def test_extract_sections():
    sections = extract_sections(SAMPLE_RESUME_TEXT)
    assert "summary" in sections
    assert "skills" in sections
    assert "experience" in sections
    assert "education" in sections

def test_parse_resume_txt_file():
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write(SAMPLE_RESUME_TEXT)
        temp_path = f.name

    try:
        raw_text, profile, skills = parse_resume_document(temp_path, "resume.txt")
        assert len(raw_text) > 100
        assert profile.contact_info.email == "jane.doe@techcorp.io"
        assert "FastAPI" in skills
        assert len(profile.experience) >= 1
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_parse_resume_docx_file():
    doc = docx.Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("jane.doe@techcorp.io | Austin, TX")
    doc.add_heading("TECHNICAL SKILLS", level=2)
    doc.add_paragraph("Python, FastAPI, Docker, PostgreSQL")
    doc.add_heading("WORK EXPERIENCE", level=2)
    doc.add_paragraph("Software Engineer | CloudScale | 2020 - Present")
    doc.add_paragraph("- Built scalable backend services in Python.")

    with tempfile.NamedTemporaryFile("wb", suffix=".docx", delete=False) as f:
        doc.save(f)
        temp_path = f.name

    try:
        raw_text, profile, skills = parse_resume_document(temp_path, "resume.docx")
        assert "FastAPI" in skills
        assert "Python" in skills
        assert profile.contact_info.email == "jane.doe@techcorp.io"
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_empty_file_error():
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as f:
        f.write("   ")
        temp_path = f.name

    try:
        with pytest.raises(ValueError, match="empty"):
            parse_resume_document(temp_path, "empty.txt")
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_unsupported_file_extension():
    with pytest.raises(ValueError, match="Unsupported file format"):
        extract_text_by_type("dummy.exe", ".exe")
